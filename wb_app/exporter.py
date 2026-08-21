from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from .calculator import calculate_scenario, distribution_status, guide_target
from .database import Database
from .models import ProductResult, RunCalculation


HEADERS = [
    "Артикул",
    "Наименование",
    "Категория",
    "Себестоимость, руб.",
    "Продажи, шт.",
    "Выручка, руб.",
    "WB реализовал, руб.",
    "К перечислению продавцу, руб.",
    "Комиссия WB (справочно), руб.",
    "Эквайринг (справочно), руб.",
    "ПВЗ (справочно), руб.",
    "Логистика, руб.",
    "Штрафы, руб.",
    "Операции на приемке, руб.",
    "Хранение, руб.",
    "Компенсация лояльности, руб.",
    "Лояльность и баллы, руб.",
    "Корректировки и удержания, руб.",
    "Прочие денежные операции, руб.",
    "Возмещение перевозчика (нейтрально), руб.",
    "Финрезультат WB, руб.",
    "Налог, руб.",
    "Себестоимость проданного, руб.",
    "Чистая прибыль, руб.",
    "Доходность",
    "Средняя розничная цена, руб.",
    "Плановая цена, руб.",
    "Плановая чистая прибыль на ед., руб.",
    "Продажи основные, шт.",
    "Продажи по выкупам, шт.",
    "Выручка основная, руб.",
    "Выручка по выкупам, руб.",
    "Цена покупателя (GMV), руб.",
    "Средняя комиссия, % от выручки",
    "Логистика, % от выручки",
    "Баллы, % от выручки",
    "Чистая прибыль, % от выручки",
]


def export_run(database: Database, run_id: int, destination: str | Path) -> Path:
    calculation = database.load_calculation(run_id)
    source_control_total = sum(
        float(row["total_amount"])
        for row in database.list_source_files(run_id)
    )
    return export_calculation(
        calculation,
        destination,
        database.planned_prices(run_id),
        source_control_total=source_control_total,
    )


def export_calculation(
    calculation: RunCalculation,
    destination: str | Path,
    planned_prices: dict[str, float] | None = None,
    source_control_total: float | None = None,
) -> Path:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Итог"
    _fill_report_sheet(
        ws,
        calculation,
        planned_prices or {},
        source_control_total=source_control_total,
    )
    _create_breakdown_sheet(workbook, calculation)
    _create_operation_guide_sheet(workbook, calculation)
    workbook.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
    output = Path(destination).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()
    return output


def suggested_export_name(calculation: RunCalculation) -> str:
    if calculation.period_start and calculation.period_end:
        period = f"{calculation.period_start:%d.%m.%Y}-{calculation.period_end:%d.%m.%Y}"
    else:
        period = datetime.now().strftime("%d.%m.%Y")
    return f"Отчет_WB_{period}.xlsx"


def _fill_report_sheet(
    ws,
    calculation: RunCalculation,
    planned_prices: dict[str, float],
    *,
    source_control_total: float | None,
) -> None:
    totals = calculation.totals()
    ws["A1"] = "WB Price Analyzer — итоговый отчет"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws["A2"] = "Период"
    ws["B2"] = (
        f"{calculation.period_start:%d.%m.%Y}–{calculation.period_end:%d.%m.%Y}"
        if calculation.period_start and calculation.period_end
        else "Не определен"
    )
    ws["D2"] = "Налоговая ставка"
    ws["E2"] = calculation.tax_rate
    ws["G2"] = "Нераспределенные доходы / расходы"
    ws["H2"] = calculation.unallocated_total
    ws["J2"] = "Чистая прибыль с нераспределенными"
    ws["K2"] = totals["net_profit"]
    channel_model = any(
        item.main_revenue is not None or item.buyout_revenue is not None
        for item in calculation.products
    )
    ws["M2"] = (
        "Контроль модели MAIN + BUYOUT"
        if channel_model
        else "Контроль исходных файлов"
    )
    source_total = (
        totals["financial_result"]
        if channel_model
        else (
            source_control_total
            if source_control_total is not None
            else sum(source.total_amount for source in calculation.source_files)
        )
    )
    ws["N2"] = source_total
    ws["P2"] = "Отклонение"
    ws["Q2"] = source_total - totals["financial_result"] if (calculation.source_files or source_control_total is not None) else 0.0

    for column, caption in enumerate(HEADERS, start=1):
        ws.cell(4, column, caption)

    for row_number, result in enumerate(calculation.products, start=5):
        _write_product_row(
            ws,
            row_number,
            result,
            calculation.tax_rate,
            planned_prices.get(result.article),
        )

    total_row = 5 + len(calculation.products)
    ws.cell(total_row, 1, "Итого по товарам")
    for column in tuple(range(5, 25)) + tuple(range(29, 34)):
        ws.cell(
            total_row,
            column,
            sum(
                float(ws.cell(row, column).value or 0)
                for row in range(5, total_row)
            ),
        )
    product_cost = float(ws.cell(total_row, 23).value or 0)
    product_net = float(ws.cell(total_row, 24).value or 0)
    product_units = float(ws.cell(total_row, 5).value or 0)
    total_profitability: float | str
    if product_units <= 0:
        total_profitability = "Нет продаж"
    elif product_cost <= 0:
        total_profitability = "Нет себестоимости"
    else:
        total_profitability = product_net / product_cost
    ws.cell(total_row, 25, total_profitability)
    ws.cell(total_row + 1, 1, "Итого с нераспределенными")
    ws.cell(total_row + 1, 21, float(ws.cell(total_row, 21).value or 0) + calculation.unallocated_total)
    ws.cell(total_row + 1, 23, product_cost)
    ws.cell(total_row + 1, 24, product_net + calculation.unallocated_total)
    ws.cell(total_row + 1, 25, total_profitability)

    revenue = float(ws.cell(total_row, 6).value or 0)
    shares = calculation.revenue_shares()
    product_net_margin = product_net / revenue if revenue else 0.0
    share_values = (
        shares["commission_share"],
        shares["logistics_share"],
        shares["points_share"],
    )
    for offset, value in enumerate(share_values, start=34):
        ws.cell(total_row, offset, value)
        ws.cell(total_row + 1, offset, value)
    ws.cell(total_row, 37, product_net_margin)
    ws.cell(total_row + 1, 37, shares["net_margin"])

    _style_sheet(ws, total_row + 1, len(HEADERS))
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(HEADERS))}{total_row - 1}"


def _write_product_row(
    ws,
    row: int,
    result: ProductResult,
    tax_rate: float,
    planned_price: float | None,
) -> None:
    scenario = calculate_scenario(result, tax_rate, planned_price)
    values = [
        result.article,
        result.name,
        result.category,
        result.total_cost,
        result.units,
        result.revenue_including_points,
        result.realized_price_total,
        result.seller_payout,
        result.wb_commission,
        result.acquiring,
        result.pvz_reimbursement,
        result.logistics_cost,
        result.penalty_cost,
        result.acceptance_cost,
        result.storage_cost,
        result.loyalty_compensation,
        result.loyalty_cost,
        result.adjustments,
        result.other,
        result.carrier_reimbursement,
        result.financial_result,
        result.tax(tax_rate),
        result.cost_sold,
        result.net_profit(tax_rate),
        _profitability_export_value(result, tax_rate),
        result.average_price(),
        scenario.planned_price,
        scenario.net_profit_per_unit,
        result.main_units_total,
        result.buyout_units_total,
        result.main_revenue_total,
        result.buyout_revenue_total,
        result.retail_price_total,
        result.commission_share(),
        result.logistics_share(),
        result.points_share(),
        result.net_margin(tax_rate),
    ]
    for column, value in enumerate(values, start=1):
        ws.cell(row, column, value)


def _profitability_export_value(result: ProductResult, tax_rate: float) -> float | str:
    if result.units <= 0:
        return "Нет продаж"
    if result.cost_sold <= 0:
        return "Нет себестоимости"
    return result.profitability(tax_rate)


def _create_breakdown_sheet(workbook, calculation: RunCalculation) -> None:
    ws = workbook.create_sheet("Разбивка")
    ws.append(["Нераспределенные доходы / расходы WB"])
    ws.append(["Тип операции", "Количество строк", "Сумма, руб."])
    for payment_reason, (row_count, amount) in calculation.unallocated.items():
        ws.append([payment_reason, row_count, amount])
    ws.append(
        [
            "Итого",
            sum(value[0] for value in calculation.unallocated.values()),
            calculation.unallocated_total,
        ]
    )
    _style_simple_table(ws, 2, ws.max_row, 3)
    ws.column_dimensions["A"].width = 72
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 20
    ws.freeze_panes = "A3"


def _create_operation_guide_sheet(workbook, calculation: RunCalculation) -> None:
    ws = workbook.create_sheet("Справочник операций")
    ws.append(["Обоснование для оплаты", "Правило", "Распределение", "С артикулом", "Без артикула"])
    for payment_reason, (with_article, without_article) in calculation.accrual_stats.items():
        ws.append(
            [
                payment_reason,
                guide_target(payment_reason),
                distribution_status(with_article, without_article),
                with_article,
                without_article,
            ]
        )
    _style_simple_table(ws, 1, ws.max_row, 5)
    for column, width in enumerate((48, 76, 34, 16, 16), start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "A2"


def _style_sheet(ws, last_row: int, last_column: int) -> None:
    dark_blue = PatternFill("solid", fgColor="1F4E78")
    light_blue = PatternFill("solid", fgColor="D9EAF7")
    light_green = PatternFill("solid", fgColor="E2F0D9")
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"].fill = dark_blue
    ws["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    for cell in ws[4]:
        cell.fill = dark_blue
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in ws.iter_rows(min_row=5, max_row=last_row, min_col=1, max_col=last_column):
        for cell in row:
            cell.border = border
    for cell in ws[last_row - 1]:
        cell.fill = light_blue
        cell.font = Font(bold=True)
    for cell in ws[last_row]:
        cell.fill = light_green
        cell.font = Font(bold=True)
    ws.row_dimensions[4].height = 46
    widths = [15, 34, 22] + [18] * (last_column - 3)
    for column, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    for row in range(2, last_row + 1):
        for column in range(4, last_column + 1):
            ws.cell(row, column).number_format = "#,##0.00"
        ws.cell(row, 25).number_format = "0.00%"
        for column in range(34, 38):
            ws.cell(row, column).number_format = "0.00%"
    ws["E2"].number_format = "0.00%"


def _style_simple_table(ws, header_row: int, last_row: int, last_column: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[header_row]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row in ws.iter_rows(min_row=header_row, max_row=last_row, min_col=1, max_col=last_column):
        for cell in row:
            cell.border = border
