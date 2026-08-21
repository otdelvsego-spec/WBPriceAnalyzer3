from __future__ import annotations

import hashlib
import math
import re
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .models import AccrualRow, BuyoutNoticeRow, ParsedSource, RealizationRow


REPORT_WEEKLY = "WEEKLY_WB"
REPORT_BUYOUT_NOTICE = "BUYOUT_NOTICE_WB"
# Compatibility constants used by inherited UI/service code.
REPORT_ACCRUAL = REPORT_WEEKLY
REPORT_REALIZATION = "UNUSED_WB_REALIZATION"

KNOWN_WB_HEADERS = {
    "№", "Номер поставки", "Предмет", "Код номенклатуры", "Бренд",
    "Артикул поставщика", "Название", "Размер", "Баркод", "Тип документа",
    "Обоснование для оплаты", "Дата заказа покупателем", "Дата продажи", "Кол-во",
    "Цена розничная", "Вайлдберриз реализовал Товар (Пр)",
    "Согласованный продуктовый дисконт, %", "Промокод, %",
    "Итоговая согласованная скидка, %", "Цена розничная с учетом согласованной скидки",
    "Размер снижения кВВ из-за рейтинга, %", "Размер изменения кВВ из-за акции, %",
    "Платформенные скидки, %", "Размер кВВ, %", "Размер кВВ без НДС, % Базовый",
    "Итоговый кВВ без НДС, %",
    "Вознаграждение с продаж до вычета услуг поверенного, без НДС",
    "Возмещение за выдачу и возврат товаров на ПВЗ",
    "Компенсация платёжных услуг/Комиссия за интеграцию платёжных сервисов",
    "Размер компенсации платёжных услуг/Комиссии за интеграцию платёжных сервисов, %",
    "Тип платежа: компенсация платёжных услуг/Комиссия за интеграцию платёжных сервисов",
    "Вознаграждение Вайлдберриз (ВВ), без НДС", "НДС с Вознаграждения Вайлдберриз",
    "К перечислению Продавцу за реализованный Товар", "Количество доставок",
    "Количество возврата", "Услуги по доставке товара покупателю",
    "Коэффициент логистики",
    "Дата начала действия фиксации", "Дата конца действия фиксации",
    "Признак услуги платной доставки", "Общая сумма штрафов",
    "Корректировка Вознаграждения Вайлдберриз (ВВ)",
    "Виды логистики, штрафов и корректировок ВВ", "Стикер МП",
    "Наименование банка-эквайера", "Номер офиса", "Наименование офиса доставки",
    "ИНН партнера", "Партнер", "Склад", "Страна", "Тип коробов",
    "Номер таможенной декларации", "Номер сборочного задания", "Код маркировки",
    "ШК", "Srid", "Возмещение издержек по перевозке/по складским операциям с товаром",
    "Организатор перевозки", "Хранение", "Удержания", "Операции на приемке",
    "Фиксированный коэффициент склада по поставке", "Признак продажи юридическому лицу",
    "Номер короба для обработки товара", "Скидка по программе софинансирования",
    "Скидка Wibes, %", "Компенсация скидки по программе лояльности",
    "Стоимость участия в программе лояльности",
    "Сумма баллов, удержанных по программе лояльности", "Id корзины заказа",
    "Разовое изменение срока перечисления денежных средств",
    "Id собственной акции продавца с дополнительной скидкой",
    "Размер дополнительной скидки по собственной акции продавца, %",
    "Способы продажи и тип товара",
    "Уникальный идентификатор скидки лояльности от продавца",
    "Размер скидки лояльности от продавца, %", "Id промокода", "Скидка за промокод, %",
    "Id подменного артикула", "Скидка по подменному артикулу, %",
    "Оптовая скидка для бизнеса, %", "ИНН покупателя-юрлица или ИП",
    "Оплата социальным сертификатом",
}


class ReportFormatError(ValueError):
    """Raised when an XLSX file is not a WB weekly detailed report."""


def normalize_text(value: object) -> str:
    text = "" if value is None else str(value)
    return " ".join(text.replace("\r", " ").replace("\n", " ").strip().split()).casefold()


def display_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_sku(value: object) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return str(int(float(value))) if float(value).is_integer() else format(float(value), "f").rstrip("0").rstrip(".")
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def as_float(value: object, default: float = 0.0) -> float:
    if value is None or value == "" or isinstance(value, bool):
        return default
    if isinstance(value, (int, float)):
        result = float(value)
        return result if math.isfinite(result) else default
    text = str(value).replace("\u00a0", "").replace(" ", "").replace(",", ".").strip()
    try:
        result = float(text)
    except ValueError:
        return default
    return result if math.isfinite(result) else default


def is_numeric(value: object) -> bool:
    if value is None or value == "" or isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return math.isfinite(float(value))
    text = str(value).replace("\u00a0", "").replace(" ", "").replace(",", ".").strip()
    try:
        return math.isfinite(float(text))
    except ValueError:
        return False


def as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None
    text = str(value).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _row_caption_map(ws, row_number: int) -> dict[str, int]:
    result: dict[str, int] = {}
    for cell in ws[row_number]:
        key = normalize_text(cell.value)
        if key and key not in result:
            result[key] = cell.column
    return result


def _find(columns: dict[str, int], *captions: str) -> int:
    for caption in captions:
        value = columns.get(normalize_text(caption), 0)
        if value:
            return value
    return 0


def _required(columns: dict[str, int], caption: str, *aliases: str) -> int:
    value = _find(columns, caption, *aliases)
    if not value:
        raise ReportFormatError(f"В отчете отсутствует обязательный столбец «{caption}»")
    return value


def _detect_weekly_sheet(workbook) -> tuple[object, int]:
    required = (
        "Артикул поставщика",
        "Обоснование для оплаты",
        "К перечислению Продавцу за реализованный Товар",
    )
    for ws in workbook.worksheets:
        for row_number in range(1, min(ws.max_row, 10) + 1):
            columns = _row_caption_map(ws, row_number)
            if all(_find(columns, caption) for caption in required):
                return ws, row_number
    raise ReportFormatError(
        "Еженедельный детализированный отчет WB не найден."
    )


def _notice_columns(ws, row_number: int) -> dict[str, int]:
    columns = _row_caption_map(ws, row_number)
    result: dict[str, int] = {}
    for caption, column in columns.items():
        if caption == normalize_text("Артикул"):
            result["article"] = column
        elif caption.startswith(normalize_text("Наименование")):
            result["name"] = column
        elif caption == normalize_text("Количество"):
            result["quantity"] = column
        elif caption.startswith(normalize_text("Сумма выкупа")):
            result["amount"] = column
    return result


def _detect_notice_sheet(workbook) -> tuple[object, int]:
    for ws in workbook.worksheets:
        for row_number in range(1, min(ws.max_row, 20) + 1):
            columns = _notice_columns(ws, row_number)
            if {"article", "quantity", "amount"}.issubset(columns):
                return ws, row_number
    raise ReportFormatError("Уведомление о выкупе WB не найдено.")


def parse_report(path: str | Path) -> ParsedSource:
    source_path = Path(path).expanduser().resolve()
    if source_path.suffix.casefold() != ".xlsx":
        raise ReportFormatError(f"Поддерживаются только файлы XLSX: {source_path.name}")
    try:
        # Normal mode is intentional: some WB exports have incomplete worksheet
        # dimension metadata and read-only mode exposes only column A.
        workbook = load_workbook(source_path, read_only=False, data_only=True)
    except Exception as exc:
        raise ReportFormatError(f"Не удалось прочитать {source_path.name}: {exc}") from exc
    try:
        try:
            ws, header_row = _detect_weekly_sheet(workbook)
            report_type = REPORT_WEEKLY
        except ReportFormatError:
            try:
                ws, header_row = _detect_notice_sheet(workbook)
                report_type = REPORT_BUYOUT_NOTICE
            except ReportFormatError as exc:
                raise ReportFormatError(
                    "Тип файла не определен. Выберите еженедельный детализированный "
                    "отчет WB или уведомление о выкупе WB в XLSX."
                ) from exc
        parsed = ParsedSource(
            path=source_path,
            file_hash=sha256_file(source_path),
            report_type=report_type,
            sheet_name=ws.title,
            header_row=header_row,
            report_number=_report_number(source_path.name),
        )
        if report_type == REPORT_BUYOUT_NOTICE:
            _parse_buyout_notice(ws, header_row, parsed)
        else:
            _parse_weekly(ws, header_row, parsed)
        return parsed
    finally:
        workbook.close()


def _parse_weekly(ws, header_row: int, parsed: ParsedSource) -> None:
    columns = _row_caption_map(ws, header_row)
    known_keys = {normalize_text(value) for value in KNOWN_WB_HEADERS}
    parsed.unknown_columns = [
        display_text(cell.value)
        for cell in ws[header_row]
        if normalize_text(cell.value) and normalize_text(cell.value) not in known_keys
    ]
    positions = {
        "subject": _required(columns, "Предмет"),
        "nm_id": _required(columns, "Код номенклатуры"),
        "article": _required(columns, "Артикул поставщика"),
        "name": _required(columns, "Название"),
        "document": _required(columns, "Тип документа"),
        "reason": _required(columns, "Обоснование для оплаты"),
        "order_date": _required(columns, "Дата заказа покупателем"),
        "sale_date": _required(columns, "Дата продажи"),
        "quantity": _required(columns, "Кол-во"),
        "retail": _required(columns, "Цена розничная"),
        "realized": _required(columns, "Вайлдберриз реализовал Товар (Пр)"),
        "pvz": _required(columns, "Возмещение за выдачу и возврат товаров на ПВЗ"),
        "acquiring": _required(columns, "Компенсация платёжных услуг/Комиссия за интеграцию платёжных сервисов"),
        "commission_no_vat": _required(columns, "Вознаграждение Вайлдберриз (ВВ), без НДС"),
        "commission_vat": _required(columns, "НДС с Вознаграждения Вайлдберриз"),
        "payout": _required(columns, "К перечислению Продавцу за реализованный Товар"),
        "logistics": _required(columns, "Услуги по доставке товара покупателю"),
        # WB introduced this optional diagnostic field during summer 2026.
        # Old reports do not contain it, while the monetary logistics column
        # already reflects the coefficient.
        "logistics_coefficient": _find(columns, "Коэффициент логистики"),
        "penalty": _required(columns, "Общая сумма штрафов"),
        "commission_adjustment": _required(columns, "Корректировка Вознаграждения Вайлдберриз (ВВ)"),
        "detail": _required(columns, "Виды логистики, штрафов и корректировок ВВ"),
        "country": _required(columns, "Страна"),
        "srid": _required(columns, "Srid"),
        "carrier": _required(columns, "Возмещение издержек по перевозке/по складским операциям с товаром"),
        "storage": _required(columns, "Хранение"),
        "deductions": _required(columns, "Удержания"),
        "acceptance": _required(columns, "Операции на приемке"),
        "loyalty_compensation": _required(columns, "Компенсация скидки по программе лояльности"),
        "loyalty_fee": _required(columns, "Стоимость участия в программе лояльности"),
        "loyalty_points": _required(columns, "Сумма баллов, удержанных по программе лояльности"),
        "payout_fee": _required(columns, "Разовое изменение срока перечисления денежных средств"),
    }

    dates: list[date] = []
    countries: set[str] = set()
    for row_number in range(header_row + 1, ws.max_row + 1):
        reason = display_text(ws.cell(row_number, positions["reason"]).value)
        if not reason:
            continue
        sale_date = as_date(ws.cell(row_number, positions["sale_date"]).value)
        order_date = as_date(ws.cell(row_number, positions["order_date"]).value)
        if sale_date:
            dates.append(sale_date)
        country = display_text(ws.cell(row_number, positions["country"]).value)
        if country:
            countries.add(country.casefold())
        parsed.accrual_rows.append(
            AccrualRow(
                source_name=parsed.path.name,
                sheet_name=ws.title,
                row_number=row_number,
                report_number=parsed.report_number,
                operation_date=order_date,
                sale_date=sale_date,
                document_type=display_text(ws.cell(row_number, positions["document"]).value),
                payment_reason=reason,
                article=display_text(ws.cell(row_number, positions["article"]).value),
                nm_id=normalize_sku(ws.cell(row_number, positions["nm_id"]).value),
                product_name=display_text(ws.cell(row_number, positions["name"]).value),
                subject=display_text(ws.cell(row_number, positions["subject"]).value),
                quantity=as_float(ws.cell(row_number, positions["quantity"]).value),
                retail_price=as_float(ws.cell(row_number, positions["retail"]).value),
                realized_price=as_float(ws.cell(row_number, positions["realized"]).value),
                seller_payout=as_float(ws.cell(row_number, positions["payout"]).value),
                logistics=as_float(ws.cell(row_number, positions["logistics"]).value),
                logistics_coefficient=(
                    as_float(ws.cell(row_number, positions["logistics_coefficient"]).value)
                    if positions["logistics_coefficient"]
                    else 0.0
                ),
                penalty=as_float(ws.cell(row_number, positions["penalty"]).value),
                storage=as_float(ws.cell(row_number, positions["storage"]).value),
                acceptance=as_float(ws.cell(row_number, positions["acceptance"]).value),
                commission_adjustment=as_float(ws.cell(row_number, positions["commission_adjustment"]).value),
                deductions=as_float(ws.cell(row_number, positions["deductions"]).value),
                loyalty_compensation=as_float(ws.cell(row_number, positions["loyalty_compensation"]).value),
                loyalty_fee=as_float(ws.cell(row_number, positions["loyalty_fee"]).value),
                loyalty_points=as_float(ws.cell(row_number, positions["loyalty_points"]).value),
                payout_fee=as_float(ws.cell(row_number, positions["payout_fee"]).value),
                acquiring=as_float(ws.cell(row_number, positions["acquiring"]).value),
                pvz_reimbursement=as_float(ws.cell(row_number, positions["pvz"]).value),
                wb_commission=(
                    as_float(ws.cell(row_number, positions["commission_no_vat"]).value)
                    + as_float(ws.cell(row_number, positions["commission_vat"]).value)
                ),
                carrier_reimbursement=as_float(ws.cell(row_number, positions["carrier"]).value),
                country=country,
                srid=display_text(ws.cell(row_number, positions["srid"]).value),
                operation_detail=display_text(ws.cell(row_number, positions["detail"]).value),
            )
        )

    if not parsed.accrual_rows:
        raise ReportFormatError(f"В отчете нет строк операций: {parsed.path.name}")

    parsed.period_start, parsed.period_end = _dominant_iso_week(dates)
    if parsed.period_start is not None and parsed.period_end is not None:
        parsed.out_of_period_rows = sum(
            1
            for row in parsed.accrual_rows
            if row.sale_date is not None
            and not (parsed.period_start <= row.sale_date <= parsed.period_end)
        )
    eaeu_non_russia = {"беларусь", "казахстан", "армения", "кыргызстан", "киргизия"}
    parsed.report_variant = (
        "по выкупам"
        if countries and countries.issubset(eaeu_non_russia)
        else "основной"
    )


def _parse_buyout_notice(ws, header_row: int, parsed: ParsedSource) -> None:
    columns = _notice_columns(ws, header_row)
    title_text = " ".join(
        display_text(cell.value)
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, header_row))
        for cell in row
        if cell.value not in (None, "")
    )
    title_match = re.search(
        r"уведомление\s+о\s+выкупе\s+№\s*(\d+)\s+от\s+(\d{4}-\d{2}-\d{2}|\d{2}\.\d{2}\.\d{4})",
        title_text,
        flags=re.IGNORECASE,
    )
    notice_date: date | None = None
    if title_match:
        parsed.report_number = title_match.group(1)
        notice_date = as_date(title_match.group(2))
    for row_number in range(header_row + 1, ws.max_row + 1):
        first_value = normalize_text(ws.cell(row_number, 1).value)
        if first_value.startswith(normalize_text("Итого")):
            break
        article = normalize_sku(ws.cell(row_number, columns["article"]).value)
        quantity_value = ws.cell(row_number, columns["quantity"]).value
        amount_value = ws.cell(row_number, columns["amount"]).value
        if not article or not is_numeric(quantity_value) or not is_numeric(amount_value):
            continue
        parsed.buyout_notice_rows.append(
            BuyoutNoticeRow(
                source_name=parsed.path.name,
                sheet_name=ws.title,
                row_number=row_number,
                report_number=parsed.report_number,
                notice_date=notice_date,
                article=article,
                product_name=(
                    display_text(ws.cell(row_number, columns["name"]).value)
                    if columns.get("name")
                    else article
                ),
                quantity=as_float(quantity_value),
                amount=as_float(amount_value),
            )
        )
    if not parsed.buyout_notice_rows:
        raise ReportFormatError(f"В уведомлении о выкупе нет товарных строк: {parsed.path.name}")
    if not parsed.report_number:
        raise ReportFormatError(
            f"Не удалось определить номер уведомления о выкупе: {parsed.path.name}"
        )
    parsed.report_variant = "уведомление о выкупе"
    if notice_date is not None:
        parsed.period_start = notice_date - timedelta(days=notice_date.weekday())
        parsed.period_end = parsed.period_start + timedelta(days=6)


def _dominant_iso_week(dates: list[date]) -> tuple[date | None, date | None]:
    if not dates:
        return None, None
    week_starts = [value - timedelta(days=value.weekday()) for value in dates]
    start = Counter(week_starts).most_common(1)[0][0]
    return start, start + timedelta(days=6)


def _report_number(filename: str) -> str:
    match = re.search(r"№\s*(\d+)", filename)
    return match.group(1) if match else ""


def workbook_sheet_names(path: str | Path) -> list[str]:
    workbook = load_workbook(Path(path), read_only=False, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def preview_sheet(
    path: str | Path,
    sheet_name: str,
    max_rows: int = 500,
    max_columns: int = 85,
) -> tuple[list[str], list[list[str]]]:
    workbook = load_workbook(Path(path), read_only=False, data_only=True)
    try:
        ws = workbook[sheet_name]
        column_count = min(ws.max_column, max_columns)
        headers = ["Строка"] + [get_column_letter(index) for index in range(1, column_count + 1)]
        rows: list[list[str]] = []
        for row_index, values in enumerate(
            ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_rows), max_col=column_count, values_only=True),
            start=1,
        ):
            rows.append([str(row_index)] + [_format_preview_value(value) for value in values])
        return headers, rows
    finally:
        workbook.close()


def _format_preview_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.6f}".rstrip("0").rstrip(".")
    return str(value)


def all_rows(sources: Iterable[ParsedSource]) -> tuple[list[AccrualRow], list[RealizationRow]]:
    rows: list[AccrualRow] = []
    for source in sources:
        rows.extend(source.accrual_rows)
    return rows, []
