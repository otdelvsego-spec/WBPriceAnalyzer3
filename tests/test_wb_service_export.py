from __future__ import annotations

import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

from wb_app.exporter import export_calculation
from wb_app.models import ParsedSource, ProductResult, RunCalculation
from wb_app.service import split_import_sources


def source(name: str, variant: str, start: date, end: date) -> ParsedSource:
    return ParsedSource(
        path=Path(name), file_hash=name, report_type="WEEKLY_WB", sheet_name="Sheet1",
        header_row=1, period_start=start, period_end=end, report_variant=variant,
    )


class WBServiceExportTests(unittest.TestCase):
    def test_groups_main_and_buyout_by_week(self) -> None:
        first = source("main.xlsx", "основной", date(2026, 8, 3), date(2026, 8, 9))
        second = source("buyout.xlsx", "по выкупам", date(2026, 8, 3), date(2026, 8, 9))
        next_week = source("next.xlsx", "основной", date(2026, 8, 10), date(2026, 8, 16))
        sessions = split_import_sources([next_week, second, first])
        self.assertEqual(len(sessions), 2)
        self.assertEqual(len(sessions[0].sources), 2)
        self.assertTrue(sessions[0].has_realization)

    def test_export_contains_wb_result_breakdown_and_guide(self) -> None:
        calculation = RunCalculation(
            run_id=None, period_start=date(2026, 8, 3), period_end=date(2026, 8, 9),
            tax_rate=0.06,
            products=[
                ProductResult(
                    "A", "Товар A", 80, 20, category="Категория", units=2,
                    revenue_no_points=2000, partner_programs=1600, points=1400,
                    delivery=-200, financial_result=1200, carrier_reimbursement=50,
                )
            ],
            unallocated_total=-50,
            unallocated={"Хранение": (1, -50)},
            accrual_stats={"Продажа": (1, 0), "Хранение": (0, 1)},
        )
        with tempfile.TemporaryDirectory() as temp_name:
            path = Path(temp_name) / "report.xlsx"
            export_calculation(calculation, path, {"A": 1200})
            workbook = load_workbook(path, data_only=False)
            self.assertEqual(workbook.sheetnames, ["Итог", "Разбивка", "Справочник операций"])
            self.assertEqual(workbook["Итог"]["A5"].value, "A")
            self.assertEqual(workbook["Итог"]["T5"].value, 50)
            self.assertEqual(workbook["Разбивка"]["C3"].value, -50)
            self.assertEqual(workbook["Справочник операций"]["A2"].value, "Продажа")
            workbook.close()


if __name__ == "__main__":
    unittest.main()
