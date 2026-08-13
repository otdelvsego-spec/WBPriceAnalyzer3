from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from wb_app.costs import (
    CostCatalogError,
    CostEditorEntry,
    build_cost_changes,
    build_products_from_editor_entries,
    export_cost_catalog,
    read_cost_catalog,
)
from wb_app.database import Database
from wb_app.models import Product


class CostCatalogTests(unittest.TestCase):
    def test_export_read_preview_and_atomic_history(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source_products = [
                Product("A-001", "Товар A", material_cost=80, labor_cost=20, active=True, category="Категория A"),
                Product("B-002", "Товар B", material_cost=45, labor_cost=5, active=False, category="Категория B"),
            ]
            path = export_cost_catalog(source_products, root / "costs.xlsx")
            imported = read_cost_catalog(path)

            self.assertEqual([product.article for product in imported], ["A-001", "B-002"])
            self.assertEqual(imported[0].total_cost, 100)
            self.assertEqual(imported[0].labor_cost, 20)
            self.assertEqual(imported[0].category, "Категория A")
            self.assertFalse(imported[1].active)

            existing = {"A-001": Product("A-001", "Товар A", material_cost=70, labor_cost=20)}
            changes = build_cost_changes(imported, existing)
            self.assertEqual([change.status for change in changes], ["Изменение", "Новая позиция"])

            database = Database(root / "app.sqlite3")
            changed = database.save_products(imported, source="Тестовый импорт")
            self.assertEqual(changed, 2)
            history = database.list_product_cost_history()
            self.assertEqual(len(history), 2)
            self.assertTrue(all(row["change_source"] == "Тестовый импорт" for row in history))

    def test_rejects_labor_above_total_and_duplicate_article(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = export_cost_catalog([], Path(directory) / "invalid.xlsx")
            workbook = load_workbook(path)
            ws = workbook["Себестоимость"]
            ws["A5"], ws["B5"], ws["D5"], ws["E5"] = "A", "Товар", 100, 120
            ws["A6"], ws["B6"], ws["D6"], ws["E6"] = "A", "Дубль", 100, 20
            ws["A7"], ws["B7"], ws["D7"], ws["E7"] = "B", "Неверные трудозатраты", 100, "нет данных"
            workbook.save(path)
            workbook.close()

            with self.assertRaises(CostCatalogError) as context:
                read_cost_catalog(path)
            self.assertIn("трудозатраты", str(context.exception))
            self.assertIn("уже указан", str(context.exception))
            self.assertIn("неверно указаны трудозатраты", str(context.exception))

    def test_builds_products_entered_inside_application(self) -> None:
        products = build_products_from_editor_entries(
            [
                CostEditorEntry("A-001", "Товар A", "125,50", "25,50", True, 1),
                CostEditorEntry("B-002", "", 80, "", False, 2, category="Категория B"),
            ]
        )

        self.assertEqual(products[0].material_cost, 100)
        self.assertEqual(products[0].labor_cost, 25.5)
        self.assertEqual(products[1].name, "B-002")
        self.assertFalse(products[1].active)
        self.assertEqual(products[1].category, "Категория B")

    def test_old_xlsx_without_category_remains_compatible(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "old_catalog.xlsx"
            workbook = Workbook()
            ws = workbook.active
            ws.title = "Себестоимость"
            ws.append(
                [
                    "Артикул",
                    "Наименование",
                    "Полная себестоимость, руб.",
                    "Трудозатраты, руб.",
                    "Активен",
                ]
            )
            ws.append(["OLD-1", "Старый товар", 120, 20, "Да"])
            workbook.save(path)
            workbook.close()

            imported = read_cost_catalog(path)
            self.assertEqual(len(imported), 1)
            self.assertEqual(imported[0].category, "")
            self.assertEqual(imported[0].total_cost, 120)

    def test_blank_total_cost_is_reported_and_skipped(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = export_cost_catalog([], Path(directory) / "incomplete.xlsx")
            workbook = load_workbook(path)
            ws = workbook["Себестоимость"]
            ws["A5"], ws["B5"], ws["D5"], ws["E5"] = "521473", "Товар", None, 10
            ws["A6"], ws["B6"], ws["D6"], ws["E6"] = "OK", "Готовый товар", 100, 20
            workbook.save(path)
            workbook.close()

            warnings: list[str] = []
            imported = read_cost_catalog(path, warnings)

            self.assertEqual([product.article for product in imported], ["OK"])
            self.assertEqual(len(warnings), 1)
            self.assertIn("521473", warnings[0])
            self.assertIn("не заполнена полная себестоимость", warnings[0])

    def test_rejects_invalid_application_entries_together(self) -> None:
        entries = [
            CostEditorEntry("A", "Товар", 100, 120, True, 1),
            CostEditorEntry("A", "Дубль", 90, 10, True, 2),
            CostEditorEntry("", "Без артикула", 50, 0, True, 3),
        ]

        with self.assertRaises(CostCatalogError) as context:
            build_products_from_editor_entries(entries)
        message = str(context.exception)
        self.assertIn("трудозатраты", message)
        self.assertIn("уже указан", message)
        self.assertIn("не указан артикул", message)


if __name__ == "__main__":
    unittest.main()
